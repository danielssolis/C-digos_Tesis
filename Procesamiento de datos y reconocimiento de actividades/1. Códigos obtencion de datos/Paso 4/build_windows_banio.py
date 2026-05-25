import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter


INPUT_FILE = "data/dataset_banio.csv"
OUTPUT_CSV = "data/dataset_banio_windows.csv"
OUTPUT_XLSX = "data/dataset_banio_windows_label.xlsx"

INPUT_SEP = ";"
OUTPUT_SEP = ";"
WINDOW_SIZE = "30s"


def auto_adjust_columns(ws, max_width=24):
    for col_cells in ws.columns:
        col_letter = get_column_letter(col_cells[0].column)
        max_length = 0

        for cell in col_cells:
            try:
                value = "" if cell.value is None else str(cell.value)
                max_length = max(max_length, len(value))
            except Exception:
                pass

        ws.column_dimensions[col_letter].width = min(max_length + 2, max_width)


def prepare_excel_for_labeling(xlsx_path: str):
    wb = load_workbook(xlsx_path)
    ws = wb["Etiquetado"]

    ws.freeze_panes = "D2"
    ws.auto_filter.ref = ws.dimensions

    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    header_map = {cell.value: cell.column for cell in ws[1]}

    activity_col = header_map.get("activity")
    completado_col = header_map.get("uso_inodoro_completado")

    label_fill = PatternFill(fill_type="solid", fgColor="FFF2CC")

    if activity_col:
        letter = get_column_letter(activity_col)
        for row in range(1, ws.max_row + 1):
            ws[f"{letter}{row}"].fill = label_fill

    if completado_col:
        letter = get_column_letter(completado_col)
        for row in range(1, ws.max_row + 1):
            ws[f"{letter}{row}"].fill = label_fill

    if "Listas" in wb.sheetnames:
        del wb["Listas"]

    ws_list = wb.create_sheet("Listas")
    ws_list["A1"] = "sin_act_banio"
    ws_list["A2"] = "uso_inodoro"
    ws_list["A3"] = "lavamanos"
    ws_list["A4"] = "ducha"

    ws_list["B1"] = 0
    ws_list["B2"] = 1
    ws_list.sheet_state = "hidden"

    if activity_col:
        letter = get_column_letter(activity_col)
        dv_activity = DataValidation(
            type="list",
            formula1="=Listas!$A$1:$A$4",
            allow_blank=True
        )
        ws.add_data_validation(dv_activity)
        dv_activity.add(f"{letter}2:{letter}{ws.max_row}")

    if completado_col:
        letter = get_column_letter(completado_col)
        dv_comp = DataValidation(
            type="list",
            formula1="=Listas!$B$1:$B$2",
            allow_blank=True
        )
        ws.add_data_validation(dv_comp)
        dv_comp.add(f"{letter}2:{letter}{ws.max_row}")

    auto_adjust_columns(ws)
    wb.save(xlsx_path)


def main():
    df = pd.read_csv(INPUT_FILE, sep=INPUT_SEP, decimal=".")
    df.columns = [str(col).strip().replace("\ufeff", "") for col in df.columns]

    print("Columnas detectadas:")
    print(df.columns.tolist())

    if "timestamp" not in df.columns:
        raise ValueError("No se encontró la columna 'timestamp'.")

    df["timestamp"] = pd.to_datetime(df["timestamp"], utc=True)

    columnas_datos = [col for col in df.columns if col != "timestamp"]
    for col in columnas_datos:
        df[col] = pd.to_numeric(df[col], errors="coerce")

    df = df.sort_values("timestamp").set_index("timestamp")
    df = df.ffill()

    agg_dict = {}

    for col in df.columns:
        c = col.lower()

        if any(x in c for x in ["temperature", "humidity", "distance"]):
            agg_dict[col] = ["mean", "min", "max", "std"]
        elif any(x in c for x in ["movement", "flow_state"]):
            agg_dict[col] = ["mean", "sum", "max"]
        else:
            agg_dict[col] = ["mean", "max"]

    df_windows = df.resample(WINDOW_SIZE).agg(agg_dict)
    df_windows.columns = [f"{col}_{stat}" for col, stat in df_windows.columns]
    df_windows = df_windows.dropna(how="all").fillna(0).reset_index()

    columnas_numericas = df_windows.select_dtypes(include=["number"]).columns
    df_windows[columnas_numericas] = df_windows[columnas_numericas].round(6)

    if "activity" not in df_windows.columns:
        df_windows["activity"] = ""

    if "uso_inodoro_completado" not in df_windows.columns:
        df_windows["uso_inodoro_completado"] = ""

    cols = df_windows.columns.tolist()
    cols_reordered = []

    for c in ["timestamp", "activity", "uso_inodoro_completado"]:
        if c in cols:
            cols_reordered.append(c)

    for c in cols:
        if c not in cols_reordered:
            cols_reordered.append(c)

    df_windows = df_windows[cols_reordered].copy()

    # Guardar CSV técnico con zona horaria
    df_windows.to_csv(
        OUTPUT_CSV,
        index=False,
        sep=OUTPUT_SEP,
        decimal=".",
        float_format="%.6f"
    )

    # Crear copia para Excel sin zona horaria
    df_excel = df_windows.copy()
    if "timestamp" in df_excel.columns:
        df_excel["timestamp"] = pd.to_datetime(
            df_excel["timestamp"], utc=True
        ).dt.tz_localize(None)

    with pd.ExcelWriter(OUTPUT_XLSX, engine="openpyxl") as writer:
        df_excel.to_excel(writer, index=False, sheet_name="Etiquetado")

    prepare_excel_for_labeling(OUTPUT_XLSX)

    print("\n✅ Dataset ventanas baño creado correctamente")
    print("CSV técnico:", OUTPUT_CSV)
    print("Excel etiquetado:", OUTPUT_XLSX)
    print("Filas:", len(df_windows))
    print("Columnas:", len(df_windows.columns))
    print("\nPrimeras filas:")
    print(df_windows.head().to_string())


if __name__ == "__main__":
    main()